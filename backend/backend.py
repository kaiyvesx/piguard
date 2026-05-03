from fastapi import FastAPI, HTTPException
from fastapi.responses import FileResponse, Response, StreamingResponse
from pydantic import BaseModel, Field

import os
import serial
import time
import threading
import re
import sqlite3
import subprocess
import glob
import cv2

from pathlib import Path
from typing import List, Dict, Optional
from datetime import datetime, timezone, timedelta

from openpyxl import Workbook
try:
    import psycopg2
    from psycopg2.extras import RealDictCursor
except Exception:
    psycopg2 = None
    RealDictCursor = None

# ===================== CONFIG =====================
UART_PORT = "/dev/serial0"
BAUD = 115200
SER_TIMEOUT = 0.3

DATA_DIR = Path.home() / "sim7000_gui_backend"
DATA_DIR.mkdir(parents=True, exist_ok=True)

DB_PATH = DATA_DIR / "app.db"
DB_ENGINE = (os.getenv("DB_ENGINE", "postgres") or "postgres").strip().lower()
if DB_ENGINE not in ("postgres", "sqlite"):
    DB_ENGINE = "postgres"

POSTGRES_DSN = (os.getenv("DATABASE_URL") or "").strip()
if not POSTGRES_DSN:
    pg_host = (os.getenv("PGHOST") or "127.0.0.1").strip()
    pg_port = (os.getenv("PGPORT") or "5432").strip()
    pg_db = (os.getenv("PGDATABASE") or "piguard_db").strip()
    pg_user = (os.getenv("PGUSER") or "piguard_user").strip()
    pg_password = (os.getenv("PGPASSWORD") or "").strip()
    if pg_password:
        POSTGRES_DSN = (
            f"host={pg_host} port={pg_port} dbname={pg_db} "
            f"user={pg_user} password={pg_password}"
        )

EXPORT_DIR = DATA_DIR / "exports"
EXPORT_DIR.mkdir(parents=True, exist_ok=True)

STATUS_POLL_SEC = 30
GPS_POLL_SEC = 30
HEARTBEAT_SEC = 60
SMS_POLL_SEC = 10

OFFLINE_AFTER_SEC = 120
DEVICE_ID = "raspi-1"

KEEP_NETWORK_STATUS = 1000
KEEP_GPS_POINTS = 5000
KEEP_ACTIVITY = 2000
KEEP_SMS_SENT = 2000
KEEP_SMS_INBOX = 1000

CAMERA_PREFERRED_WIDTH = 640
CAMERA_PREFERRED_HEIGHT = 480
CAMERA_PREFERRED_FPS = 15
CAMERA_STREAM_JPEG_QUALITY = 80
# =================================================
QMI_MODEM = "/dev/cdc-wdm0"
WWAN_IFACE = "wwan0"

PH_TZ = timezone(timedelta(hours=8))


def now_dt_ph() -> datetime:
    return datetime.now(PH_TZ)


def now_iso_ph() -> str:
    return now_dt_ph().isoformat(timespec="seconds")


def excel_ts(dt_or_iso):
    if isinstance(dt_or_iso, datetime):
        dt = dt_or_iso
    else:
        try:
            dt = datetime.fromisoformat(str(dt_or_iso))
        except Exception:
            return dt_or_iso
    return dt.replace(tzinfo=None)


def make_id() -> str:
    return hex(int(time.time() * 1_000_000))[2:]


def ago_str_from_seconds(sec: Optional[int]) -> str:
    if sec is None:
        return "unknown"
    if sec < 60:
        return f"{sec}s ago"
    m = sec // 60
    if m < 60:
        return f"{m}min ago"
    h = m // 60
    if h < 24:
        return f"{h}hr ago"
    d = h // 24
    return f"{d}day(s) ago"


app = FastAPI(title="SIM7000C Backend", version="7.4")

# ===================== Serial =====================
ser = None
serial_lock = threading.Lock()
gps_power_lock = threading.Lock()
camera_lock = threading.Lock()
camera_state_lock = threading.Lock()
camera_state: Dict[int, dict] = {}


def _camera_state_get(idx: int) -> dict:
    with camera_state_lock:
        return dict(camera_state.get(idx, {}))


def _camera_state_update(idx: int, **kwargs):
    with camera_state_lock:
        st = camera_state.setdefault(idx, {})
        st.update(kwargs)



def ensure_serial():
    global ser
    if ser is None or not ser.is_open:
        ser = serial.Serial(UART_PORT, BAUD, timeout=SER_TIMEOUT)
        time.sleep(0.6)
        ser.reset_input_buffer()


def read_available() -> str:
    if ser and ser.in_waiting:
        return ser.read(ser.in_waiting).decode(errors="ignore")
    return ""


def send_at(cmd: str, max_wait: float = 1.2) -> str:
    with serial_lock:
        ensure_serial()
        ser.reset_input_buffer()
        ser.write((cmd + "\r").encode())

        t0 = time.time()
        buf = ""
        while time.time() - t0 < max_wait:
            buf += read_available()
            if "\r\nOK\r\n" in buf or "\r\nERROR\r\n" in buf:
                break
            time.sleep(0.05)
        return buf.strip()


def send_sms_single(number: str, message: str) -> str:
    with serial_lock:
        ensure_serial()
        ser.reset_input_buffer()

        ser.write(b"AT+CMGF=1\r")
        time.sleep(0.3)
        _ = read_available()

        ser.write(f'AT+CMGS="{number}"\r'.encode())
        time.sleep(0.8)
        prompt = read_available()

        ser.write(message.encode(errors="ignore"))
        ser.write(b"\x1A")

        t0 = time.time()
        buf = prompt
        while time.time() - t0 < 20:
            buf += read_available()
            if ("+CMGS" in buf and "OK" in buf) or ("ERROR" in buf):
                break
            time.sleep(0.1)
        return buf.strip()


# ===================== Helpers =====================
def normalize_number(n: str) -> str:
    s = (n or "").strip().replace(" ", "")
    if s.startswith("+63") and len(s) >= 13:
        s = "0" + s[3:]
    return s


def is_valid_ph(n: str) -> bool:
    return bool(re.fullmatch(r"09\d{9}", n))


def parse_apn_from_cgdcont(resp: str) -> Optional[str]:
    m = re.search(r'\+CGDCONT:\s*\d+,"[^"]*","([^"]+)"', resp)
    return m.group(1) if m else None


def interpret_csq(csq_resp: str) -> str:
    try:
        if "+CSQ:" not in csq_resp:
            return "Unknown signal"
        val = int(csq_resp.split(":")[1].split(",")[0].strip())
        if val == 99:
            return "Unknown signal"
        if val < 10:
            return "Very weak signal"
        if val < 15:
            return "Weak signal"
        if val < 20:
            return "Good signal"
        return "Excellent signal"
    except Exception:
        return "Unknown signal"


def interpret_creg(creg_resp: str) -> str:
    try:
        m = re.search(r"\+CREG:\s*\d+,(\d+)", creg_resp)
        if not m:
            return "Unknown network state"
        stat = m.group(1)
        if stat == "1":
            return "Registered to network"
        if stat == "0":
            return "Not registered"
        if stat == "2":
            return "Searching network"
        if stat == "5":
            return "Roaming network"
        return "Unknown network state"
    except Exception:
        return "Unknown network state"


def interpret_at(at_resp: str) -> str:
    return "OK" if "OK" in (at_resp or "") else "ERROR/NO RESPONSE"


def get_throttled_status() -> dict:
    try:
        res = subprocess.run(
            ["vcgencmd", "get_throttled"],
            capture_output=True,
            text=True,
            timeout=2
        )
        raw = (res.stdout or "").strip()

        value = 0
        if "throttled=" in raw:
            value = int(raw.split("=", 1)[1], 16)

        return {
            "raw": raw,
            "value": value,
            "undervoltage_now": 1 if (value & (1 << 0)) else 0,
            "throttling_now": 1 if (value & (1 << 2)) else 0,
            "undervoltage_boot": 1 if (value & (1 << 16)) else 0,
            "throttling_boot": 1 if (value & (1 << 18)) else 0,
        }
    except Exception as e:
        return {
            "raw": f"ERROR: {e}",
            "value": None,
            "undervoltage_now": None,
            "throttling_now": None,
            "undervoltage_boot": None,
            "throttling_boot": None,
        }


def delayed_poweroff():
    time.sleep(2)
    subprocess.run(["sudo", "poweroff"])


def run_cmd(cmd: List[str], timeout: int = 8) -> dict:
    try:
        res = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
        return {
            "ok": res.returncode == 0,
            "stdout": (res.stdout or "").strip(),
            "stderr": (res.stderr or "").strip(),
            "returncode": res.returncode,
        }
    except Exception as e:
        return {
            "ok": False,
            "stdout": "",
            "stderr": str(e),
            "returncode": -1,
        }


def parse_qmi_settings(text: str) -> dict:
    out = {
        "ip_family": None,
        "ipv4": None,
        "subnet_mask": None,
        "gateway": None,
        "dns1": None,
        "dns2": None,
        "mtu": None,
    }
    if not text:
        return out

    patterns = {
        "ip_family": r"IP Family:\s*([A-Za-z0-9]+)",
        "ipv4": r"IPv4 address:\s*([0-9.]+)",
        "subnet_mask": r"IPv4 subnet mask:\s*([0-9.]+)",
        "gateway": r"IPv4 gateway address:\s*([0-9.]+)",
        "dns1": r"IPv4 primary DNS:\s*([0-9.]+)",
        "dns2": r"IPv4 secondary DNS:\s*([0-9.]+)",
        "mtu": r"MTU:\s*([0-9]+)",
    }
    for k, pat in patterns.items():
        m = re.search(pat, text)
        if m:
            out[k] = int(m.group(1)) if k == "mtu" else m.group(1)
    return out


def parse_serving_system(text: str) -> dict:
    out = {
        "registration_state": None,
        "ps": None,
        "selected_network": None,
        "radio_interface": None,
        "data_capability": None,
        "operator": None,
    }
    if not text:
        return out

    pairs = {
        "registration_state": r"Registration state:\s*'([^']+)'",
        "ps": r"PS:\s*'([^']+)'",
        "selected_network": r"Selected network:\s*'([^']+)'",
        "operator": r"Description:\s*'([^']+)'",
    }
    for k, pat in pairs.items():
        m = re.search(pat, text)
        if m:
            out[k] = m.group(1)

    m = re.search(r"Radio interfaces:\s*'.*?'\s*\n\s*\[\d+\]:\s*'([^']+)'", text)
    if m:
        out["radio_interface"] = m.group(1)
    m = re.search(r"Data service capabilities:\s*'.*?'\s*\n\s*\[\d+\]:\s*'([^']+)'", text)
    if m:
        out["data_capability"] = m.group(1)
    return out


def parse_signal_info(text: str) -> dict:
    out = {"rssi_dbm": None, "raw": text or ""}
    if not text:
        return out
    m = re.search(r"RSSI:\s*'(-?\d+)\s*dBm'", text)
    if m:
        out["rssi_dbm"] = int(m.group(1))
    return out


def parse_ip_addr_show(text: str) -> dict:
    out = {"is_up": False, "ipv4_cidr": None, "link_state_text": text or ""}
    if not text:
        return out
    first_line = text.splitlines()[0] if text.splitlines() else ""
    out["is_up"] = ("UP" in first_line or "LOWER_UP" in first_line)
    m = re.search(r"\binet\s+([0-9.]+/\d+)", text)
    if m:
        out["ipv4_cidr"] = m.group(1)
    return out


def get_mobile_data_status() -> dict:
    result = {
        "modem_path": QMI_MODEM,
        "iface": WWAN_IFACE,
        "modem_present": Path(QMI_MODEM).exists(),
        "iface_present": Path(f"/sys/class/net/{WWAN_IFACE}").exists(),
        "connected": False,
        "session_status": "unknown",
        "wwan0_up": False,
        "wwan0_ipv4": None,
        "ip_family": None,
        "ipv4": None,
        "subnet_mask": None,
        "gateway": None,
        "dns1": None,
        "dns2": None,
        "mtu": None,
        "registration_state": None,
        "ps": None,
        "selected_network": None,
        "radio_interface": None,
        "data_capability": None,
        "operator": None,
        "rssi_dbm": None,
        "notes": [],
        "raw": {},
    }

    if not result["modem_present"]:
        result["notes"].append("QMI modem not found")
        return result

    status_cmd = run_cmd(["qmicli", "-d", QMI_MODEM, "--wds-get-packet-service-status"], timeout=8)
    result["raw"]["packet_service_status"] = status_cmd["stdout"] or status_cmd["stderr"]
    if status_cmd["ok"]:
        m = re.search(r"Connection status:\s*'([^']+)'", status_cmd["stdout"])
        if m:
            result["session_status"] = m.group(1)
            result["connected"] = (m.group(1).lower() == "connected")
    else:
        result["notes"].append("Failed to read packet service status")

    ip_show = run_cmd(["ip", "addr", "show", WWAN_IFACE], timeout=4)
    result["raw"]["ip_addr_show"] = ip_show["stdout"] or ip_show["stderr"]
    if ip_show["ok"]:
        parsed_if = parse_ip_addr_show(ip_show["stdout"])
        result["wwan0_up"] = parsed_if["is_up"]
        result["wwan0_ipv4"] = parsed_if["ipv4_cidr"]

    serving = run_cmd(["qmicli", "-d", QMI_MODEM, "--nas-get-serving-system"], timeout=8)
    result["raw"]["serving_system"] = serving["stdout"] or serving["stderr"]
    if serving["ok"]:
        result.update(parse_serving_system(serving["stdout"]))

    signal = run_cmd(["qmicli", "-d", QMI_MODEM, "--nas-get-signal-info"], timeout=8)
    result["raw"]["signal_info"] = signal["stdout"] or signal["stderr"]
    if signal["ok"]:
        result["rssi_dbm"] = parse_signal_info(signal["stdout"])["rssi_dbm"]

    if result["connected"]:
        current = run_cmd(["qmicli", "-p", "-d", QMI_MODEM, "--wds-get-current-settings"], timeout=8)
        result["raw"]["current_settings"] = current["stdout"] or current["stderr"]
        if current["ok"]:
            result.update(parse_qmi_settings(current["stdout"]))
        else:
            result["notes"].append("Connected, but current settings could not be read")

    if result["connected"] and not result["wwan0_ipv4"]:
        result["notes"].append("Packet session connected, but wwan0 has no IPv4 assigned")
    if not result["connected"]:
        result["notes"].append("Mobile data session is disconnected")
    return result


# ===================== GPS control =====================
def query_gps_power_state() -> Optional[int]:
    try:
        resp = send_at("AT+CGNSPWR?", max_wait=1.0)
        m = re.search(r"\+CGNSPWR:\s*(\d+)", resp)
        if m:
            return int(m.group(1))
    except Exception:
        pass
    return None


def gps_power_cycle_sequence(log_it: bool = True) -> dict:
    with gps_power_lock:
        try:
            _ = send_at("AT", max_wait=0.8)
            off_resp = send_at("AT+CGNSPWR=0", max_wait=1.0)
            time.sleep(0.8)
            on_resp = send_at("AT+CGNSPWR=1", max_wait=1.2)
            time.sleep(2.0)
            state = query_gps_power_state()

            result = {
                "ok": True,
                "off_resp": off_resp,
                "on_resp": on_resp,
                "power_state": state,
            }
            if log_it:
                log_activity("INFO", "GPS", f"gps_power_cycle: power_state={state}")
            return result
        except Exception as e:
            if log_it:
                log_activity("ERROR", "GPS", f"gps_power_cycle: {e}")
            return {"ok": False, "error": str(e)}


def ensure_gps_on() -> bool:
    state = query_gps_power_state()
    if state == 1:
        return True

    result = gps_power_cycle_sequence(log_it=True)
    return bool(result.get("ok"))


# ===================== SMS inbox =====================
def configure_sms_mode():
    with serial_lock:
        ensure_serial()
        ser.reset_input_buffer()

        ser.write(b"AT+CMGF=1\r")
        time.sleep(0.3)
        _ = read_available()

        ser.write(b"AT+CNMI=2,1,0,0,0\r")
        time.sleep(0.3)
        _ = read_available()


def parse_cmgl_text(resp: str) -> List[dict]:
    out = []
    if not resp:
        return out

    lines = resp.replace("\r\n", "\n").split("\n")
    i = 0

    while i < len(lines):
        line = lines[i].strip()

        if not line.startswith("+CMGL:"):
            i += 1
            continue

        header = line
        message_lines = []
        i += 1

        while i < len(lines):
            nxt = lines[i]
            stripped = nxt.strip()
            if stripped.startswith("+CMGL:") or stripped == "OK" or stripped == "ERROR":
                break
            message_lines.append(nxt)
            i += 1

        message = "\n".join(message_lines).strip()

        m = re.match(
            r'^\+CMGL:\s*(\d+),"([^"]*)","([^"]*)",(?:"([^"]*)"|),"?([^"]*)"?$',
            header
        )

        if not m:
            continue

        modem_index = int(m.group(1))
        sms_status = (m.group(2) or "").strip()
        from_number = normalize_number((m.group(3) or "").strip())
        sender_ts = (m.group(5) or "").strip()

        raw_block = header
        if message:
            raw_block += "\n" + message

        out.append({
            "modem_index": modem_index,
            "sms_status": sms_status,
            "from_number": from_number,
            "sender_ts": sender_ts,
            "message": message,
            "raw": raw_block.strip(),
        })

    return out


def sms_exists(raw: str) -> bool:
    rows = db_query("SELECT id FROM sms_inbox WHERE raw=? LIMIT 1", (raw,))
    return bool(rows)


def fetch_inbox_once() -> List[dict]:
    configure_sms_mode()

    resp = send_at('AT+CMGL="ALL",1', max_wait=6.0)
    msgs = parse_cmgl_text(resp)
    if msgs:
        return msgs

    try:
        send_at('AT+CPMS="SM","SM","SM"', max_wait=2.0)
        resp = send_at('AT+CMGL="ALL",1', max_wait=6.0)
        msgs = parse_cmgl_text(resp)
        if msgs:
            return msgs
    except Exception:
        pass

    try:
        send_at('AT+CPMS="ME","ME","ME"', max_wait=2.0)
        resp = send_at('AT+CMGL="ALL",1', max_wait=6.0)
        msgs = parse_cmgl_text(resp)
        if msgs:
            return msgs
    except Exception:
        pass

    return []


def save_incoming_sms(msg: dict):
    raw = msg.get("raw", "")
    if not raw or sms_exists(raw):
        return

    from_number = normalize_number(msg.get("from_number", "")) or "UNKNOWN"
    message = msg.get("message", "")

    db_exec(
        """
        INSERT INTO sms_inbox (id, ts, from_number, message, raw)
        VALUES (?, ?, ?, ?, ?)
        """,
        (
            make_id(),
            now_iso_ph(),
            from_number,
            message,
            raw,
        )
    )
    db_prune_keep_latest("sms_inbox", KEEP_SMS_INBOX)

    log_activity(
        "INFO",
        "SMS",
        f"inbox_receive: from={from_number} modem_index={msg.get('modem_index')} status={msg.get('sms_status')}"
    )


def cleanup_modem_sms(modem_index: int):
    try:
        send_at(f"AT+CMGD={int(modem_index)}", max_wait=1.5)
    except Exception as e:
        log_activity("WARN", "SMS", f"cleanup_modem_sms index={modem_index}: {e}")


def sms_inbox_poll_once():
    msgs = fetch_inbox_once()

    if not msgs:
        try:
            raw = send_at('AT+CMGL="ALL",1', max_wait=6.0)
            log_activity("INFO", "SMS", f"sms_poll_empty: raw={raw[:250]}")
        except Exception as e:
            log_activity("ERROR", "SMS", f"sms_poll_raw_error: {e}")
        return

    for msg in msgs:
        save_incoming_sms(msg)
        if msg.get("modem_index") is not None:
            cleanup_modem_sms(msg["modem_index"])


# ===================== Camera helpers =====================
def list_video_devices() -> List[str]:
    devs = []
    for path in glob.glob("/dev/video*"):
        m = re.fullmatch(r"/dev/video(\d+)", path)
        if not m:
            continue
        idx = int(m.group(1))
        if idx <= 9:
            devs.append((idx, path))

    devs.sort(key=lambda x: x[0])
    return [path for _, path in devs]


def camera_index_from_path(dev_path: str) -> Optional[int]:
    m = re.search(r"/dev/video(\d+)$", dev_path or "")
    if not m:
        return None
    return int(m.group(1))



def camera_name_from_path(dev_path: str) -> str:
    p = Path(dev_path)
    if not p.name.startswith("video"):
        return "-"
    sys_name = Path("/sys/class/video4linux") / p.name / "name"
    try:
        return sys_name.read_text().strip() or "-"
    except Exception:
        return "-"


def probe_camera(dev_path: str) -> dict:
    idx = camera_index_from_path(dev_path)
    present = Path(dev_path).exists()
    result = {
        "device": dev_path,
        "index": idx,
        "present": present,
        "exists": present,
        "name": camera_name_from_path(dev_path),
        "ok": False,
        "busy": False,
        "streaming": False,
        "status": "UNKNOWN",
        "width": None,
        "height": None,
        "fps": None,
        "last_frame_ts": None,
        "last_error": None,
    }

    if idx is None:
        result["status"] = "INVALID"
        return result

    st = _camera_state_get(idx)
    if st:
        result.update({
            "streaming": bool(st.get("streaming", False)),
            "last_frame_ts": st.get("last_frame_ts"),
            "last_error": st.get("last_error"),
            "width": st.get("width") or result["width"],
            "height": st.get("height") or result["height"],
            "fps": st.get("fps") or result["fps"],
        })

    if result["streaming"]:
        result["ok"] = True
        result["busy"] = True
        result["status"] = "LIVE"
        return result

    if not present:
        result["status"] = "MISSING"
        return result

    with camera_lock:
        cap = cv2.VideoCapture(idx, cv2.CAP_V4L2)
        try:
            if not cap.isOpened():
                result["status"] = "DETECTED_NOT_OPENED"
                return result

            cap.set(cv2.CAP_PROP_FRAME_WIDTH, CAMERA_PREFERRED_WIDTH)
            cap.set(cv2.CAP_PROP_FRAME_HEIGHT, CAMERA_PREFERRED_HEIGHT)
            cap.set(cv2.CAP_PROP_FPS, CAMERA_PREFERRED_FPS)
            cap.set(cv2.CAP_PROP_BUFFERSIZE, 1)

            ok = False
            frame = None
            for _ in range(3):
                ok, frame = cap.read()
                if ok and frame is not None:
                    break
                time.sleep(0.03)

            if not ok or frame is None:
                result["status"] = "OPEN_BUT_NO_FRAME"
                return result

            h, w = frame.shape[:2]
            fps = cap.get(cv2.CAP_PROP_FPS)

            result.update({
                "ok": True,
                "status": "READY",
                "width": int(w),
                "height": int(h),
                "fps": float(fps) if fps else None,
            })
            _camera_state_update(
                idx,
                width=int(w),
                height=int(h),
                fps=float(fps) if fps else None,
                last_error=None
            )
            return result
        finally:
            cap.release()


def list_cameras() -> List[dict]:
    cameras = []
    for dev in list_video_devices():
        cameras.append(probe_camera(dev))
    return cameras

def capture_camera_jpeg(camera_index: int) -> bytes:
    with camera_lock:
        cap = cv2.VideoCapture(camera_index, cv2.CAP_V4L2)
        if not cap.isOpened():
            _camera_state_update(camera_index, last_error=f"/dev/video{camera_index} not available")
            raise RuntimeError(f"Camera /dev/video{camera_index} not available")

        try:
            cap.set(cv2.CAP_PROP_FRAME_WIDTH, CAMERA_PREFERRED_WIDTH)
            cap.set(cv2.CAP_PROP_FRAME_HEIGHT, CAMERA_PREFERRED_HEIGHT)
            cap.set(cv2.CAP_PROP_FPS, CAMERA_PREFERRED_FPS)
            cap.set(cv2.CAP_PROP_BUFFERSIZE, 1)

            ok = False
            frame = None
            for _ in range(3):
                ok, frame = cap.read()
                if ok and frame is not None:
                    break
                time.sleep(0.03)

            if not ok or frame is None:
                _camera_state_update(camera_index, last_error="No frame captured in snapshot")
                raise RuntimeError(f"Camera /dev/video{camera_index} not available")

            h, w = frame.shape[:2]
            fps = cap.get(cv2.CAP_PROP_FPS)
            _camera_state_update(camera_index, width=int(w), height=int(h), fps=float(fps) if fps else None, last_error=None)

            ok2, buf = cv2.imencode(".jpg", frame)
            if not ok2:
                _camera_state_update(camera_index, last_error="Failed to encode JPEG snapshot")
                raise RuntimeError("Failed to encode JPEG")

            return buf.tobytes()
        finally:
            cap.release()


def mjpeg_stream_generator(camera_index: int):
    _camera_state_update(camera_index, streaming=True, status="LIVE", last_error=None)
    cap = cv2.VideoCapture(camera_index, cv2.CAP_V4L2)
    if not cap.isOpened():
        _camera_state_update(camera_index, streaming=False, status="ERROR", last_error=f"Camera /dev/video{camera_index} not available")
        raise RuntimeError(f"Camera /dev/video{camera_index} not available")

    try:
        cap.set(cv2.CAP_PROP_FRAME_WIDTH, CAMERA_PREFERRED_WIDTH)
        cap.set(cv2.CAP_PROP_FRAME_HEIGHT, CAMERA_PREFERRED_HEIGHT)
        cap.set(cv2.CAP_PROP_FPS, CAMERA_PREFERRED_FPS)
        cap.set(cv2.CAP_PROP_BUFFERSIZE, 1)

        encode_param = [int(cv2.IMWRITE_JPEG_QUALITY), CAMERA_STREAM_JPEG_QUALITY]

        while True:
            ok, frame = cap.read()
            if not ok or frame is None:
                _camera_state_update(camera_index, last_error="Live stream frame read failed")
                time.sleep(0.05)
                continue

            h, w = frame.shape[:2]
            fps = cap.get(cv2.CAP_PROP_FPS)
            _camera_state_update(
                camera_index,
                streaming=True,
                status="LIVE",
                width=int(w),
                height=int(h),
                fps=float(fps) if fps else None,
                last_frame_ts=now_dt_ph().strftime("%m/%d/%Y %I:%M:%S %p"),
                last_error=None,
            )

            ok2, buf = cv2.imencode(".jpg", frame, encode_param)
            if not ok2:
                _camera_state_update(camera_index, last_error="JPEG encode failed during live stream")
                continue

            jpg_bytes = buf.tobytes()

            yield (
                b"--frame\r\n"
                b"Content-Type: image/jpeg\r\n"
                b"Content-Length: " + str(len(jpg_bytes)).encode() + b"\r\n\r\n"
                + jpg_bytes + b"\r\n"
            )
    finally:
        cap.release()
        _camera_state_update(camera_index, streaming=False, status="STOPPED")


# ===================== DB =====================
db_lock = threading.Lock()
_db_backend = None


def _resolve_db_backend() -> str:
    global _db_backend
    if _db_backend:
        return _db_backend

    if DB_ENGINE == "postgres":
        if psycopg2 is None:
            print("DB backend: psycopg2 not installed, falling back to sqlite")
            _db_backend = "sqlite"
            return _db_backend
        if not POSTGRES_DSN:
            print("DB backend: postgres selected but DSN not configured, falling back to sqlite")
            _db_backend = "sqlite"
            return _db_backend
        try:
            test_conn = psycopg2.connect(POSTGRES_DSN, connect_timeout=4)
            test_conn.close()
            _db_backend = "postgres"
            print("DB backend: postgres")
            return _db_backend
        except Exception as e:
            print(f"DB backend: postgres unavailable ({e}), falling back to sqlite")
            _db_backend = "sqlite"
            return _db_backend

    _db_backend = "sqlite"
    print("DB backend: sqlite")
    return _db_backend


def _sql_params(sql: str, params: tuple) -> tuple[str, tuple]:
    if _resolve_db_backend() == "postgres":
        return sql.replace("?", "%s"), params
    return sql, params


def db_connect():
    if _resolve_db_backend() == "postgres":
        return psycopg2.connect(POSTGRES_DSN)

    conn = sqlite3.connect(str(DB_PATH), check_same_thread=False)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA journal_mode=DELETE;")
    conn.execute("PRAGMA synchronous=NORMAL;")
    return conn


def db_exec(sql: str, params: tuple = ()):
    sql2, params2 = _sql_params(sql, params)
    with db_lock:
        conn = db_connect()
        try:
            if _resolve_db_backend() == "postgres":
                with conn.cursor() as cur:
                    cur.execute(sql2, params2)
            else:
                conn.execute(sql2, params2)
            conn.commit()
        finally:
            conn.close()


def db_query(sql: str, params: tuple = ()) -> List[dict]:
    sql2, params2 = _sql_params(sql, params)
    with db_lock:
        conn = db_connect()
        try:
            if _resolve_db_backend() == "postgres":
                with conn.cursor(cursor_factory=RealDictCursor) as cur:
                    cur.execute(sql2, params2)
                    rows = cur.fetchall()
                    return [dict(r) for r in rows]
            cur = conn.execute(sql2, params2)
            rows = cur.fetchall()
            return [dict(r) for r in rows]
        finally:
            conn.close()


def db_prune_keep_latest(table: str, keep: int):
    if keep <= 0:
        return
    db_exec(
        f"""
        DELETE FROM {table}
        WHERE id NOT IN (
            SELECT id FROM {table}
            ORDER BY ts DESC
            LIMIT ?
        )
        """,
        (keep,)
    )


def init_db():
    db_exec("""
    CREATE TABLE IF NOT EXISTS contacts (
        id TEXT PRIMARY KEY,
        name TEXT NOT NULL,
        number TEXT NOT NULL,
        created_at TEXT NOT NULL
    )""")

    db_exec("""
    CREATE TABLE IF NOT EXISTS groups (
        name TEXT PRIMARY KEY
    )""")

    db_exec("""
    CREATE TABLE IF NOT EXISTS group_members (
        group_name TEXT NOT NULL,
        contact_id TEXT NOT NULL,
        PRIMARY KEY (group_name, contact_id)
    )""")

    db_exec("""
    CREATE TABLE IF NOT EXISTS sms_inbox (
        id TEXT PRIMARY KEY,
        ts TEXT NOT NULL,
        from_number TEXT NOT NULL,
        message TEXT NOT NULL,
        raw TEXT
    )""")

    db_exec("""
    CREATE TABLE IF NOT EXISTS sms_sent (
        id TEXT PRIMARY KEY,
        ts TEXT NOT NULL,
        to_number TEXT NOT NULL,
        message TEXT NOT NULL,
        resp TEXT,
        ok INTEGER NOT NULL
    )""")

    db_exec("""
    CREATE TABLE IF NOT EXISTS gps_points (
        id TEXT PRIMARY KEY,
        ts TEXT NOT NULL,
        utc TEXT,
        lat REAL,
        lon REAL,
        alt REAL,
        speed REAL,
        course REAL,
        sat INTEGER,
        google_maps TEXT
    )""")

    db_exec("""
    CREATE TABLE IF NOT EXISTS network_status (
        id TEXT PRIMARY KEY,
        ts TEXT NOT NULL,
        at TEXT,
        csq TEXT,
        creg TEXT,
        operator TEXT,
        ip TEXT,
        apn TEXT,
        undervoltage_now INTEGER,
        undervoltage_boot INTEGER,
        throttling_now INTEGER,
        throttling_boot INTEGER,
        throttled_raw TEXT
    )""")

    db_exec("""
    CREATE TABLE IF NOT EXISTS heartbeat (
        device_id TEXT PRIMARY KEY,
        last_seen_ts TEXT NOT NULL,
        last_online_ts TEXT NOT NULL,
        last_offline_ts TEXT
    )""")

    db_exec("""
    CREATE TABLE IF NOT EXISTS activity (
        id TEXT PRIMARY KEY,
        ts TEXT NOT NULL,
        level TEXT NOT NULL,
        category TEXT NOT NULL,
        detail TEXT NOT NULL
    )""")


def db_reset_storage():
    if _resolve_db_backend() == "sqlite":
        if DB_PATH.exists():
            DB_PATH.unlink()
        return

    for table in [
        "group_members",
        "sms_inbox",
        "sms_sent",
        "gps_points",
        "network_status",
        "heartbeat",
        "activity",
        "contacts",
        "groups",
    ]:
        db_exec(f"DROP TABLE IF EXISTS {table}")


def log_activity(level: str, category: str, detail: str):
    try:
        db_exec(
            "INSERT INTO activity (id, ts, level, category, detail) VALUES (?, ?, ?, ?, ?)",
            (make_id(), now_iso_ph(), level, category, detail)
        )
        db_prune_keep_latest("activity", KEEP_ACTIVITY)
    except Exception as e:
        print("log_activity failed:", e)


# ===================== Heartbeat =====================
def hb_touch_online():
    ts = now_iso_ph()
    rows = db_query("SELECT device_id FROM heartbeat WHERE device_id=?", (DEVICE_ID,))
    if rows:
        db_exec(
            """
            UPDATE heartbeat
            SET last_seen_ts=?, last_online_ts=?
            WHERE device_id=?
            """,
            (ts, ts, DEVICE_ID)
        )
    else:
        db_exec(
            """
            INSERT INTO heartbeat (device_id, last_seen_ts, last_online_ts, last_offline_ts)
            VALUES (?, ?, ?, NULL)
            """,
            (DEVICE_ID, ts, ts)
        )


def hb_get_status() -> dict:
    rows = db_query("SELECT * FROM heartbeat WHERE device_id=?", (DEVICE_ID,))
    if not rows:
        return {
            "device_id": DEVICE_ID,
            "status": "OFFLINE",
            "last_seen_ts": None,
            "last_online_ts": None,
            "last_offline_ts": None,
            "last_online_ago": "unknown",
            "offline_after_sec": OFFLINE_AFTER_SEC,
        }

    r = rows[0]
    last_seen = r["last_seen_ts"]
    last_online = r["last_online_ts"]
    last_offline = r["last_offline_ts"]

    now = now_dt_ph()
    age_seen = None
    age_online = None
    try:
        age_seen = int((now - datetime.fromisoformat(last_seen)).total_seconds())
    except Exception:
        pass
    try:
        age_online = int((now - datetime.fromisoformat(last_online)).total_seconds())
    except Exception:
        pass

    status = "ONLINE" if (age_seen is not None and age_seen <= OFFLINE_AFTER_SEC) else "OFFLINE"

    if status == "OFFLINE" and not last_offline:
        try:
            ts = now_iso_ph()
            db_exec("UPDATE heartbeat SET last_offline_ts=? WHERE device_id=?", (ts, DEVICE_ID))
            last_offline = ts
        except Exception:
            pass

    return {
        "device_id": DEVICE_ID,
        "status": status,
        "last_seen_ts": last_seen,
        "last_online_ts": last_online,
        "last_offline_ts": last_offline,
        "last_online_ago": ago_str_from_seconds(age_online),
        "offline_after_sec": OFFLINE_AFTER_SEC,
    }


# ===================== GPS =====================
def parse_cgnsinf(resp: str) -> Dict:
    raw = resp
    m = re.search(r"\+CGNSINF:\s*(.+)", resp)
    if not m:
        return {"raw": raw, "run": None, "fix": None}

    parts = m.group(1).split(",")
    parts += [""] * (25 - len(parts))

    run = parts[0].strip() if parts[0] else None
    fix = parts[1].strip() if parts[1] else None
    utc = parts[2].strip() if parts[2] else None

    def to_float(s):
        try:
            return float(s) if s not in (None, "", "0") else None
        except Exception:
            return None

    lat = to_float(parts[3].strip()) if parts[3] else None
    lon = to_float(parts[4].strip()) if parts[4] else None
    alt = to_float(parts[5].strip()) if parts[5] else None
    speed = to_float(parts[6].strip()) if parts[6] else None
    course = to_float(parts[7].strip()) if parts[7] else None

    sat = None
    try:
        sat = int(parts[14].strip()) if parts[14] else None
    except Exception:
        sat = None

    out = {
        "raw": raw,
        "run": run,
        "fix": fix,
        "utc": utc,
        "lat": lat,
        "lon": lon,
        "alt": alt,
        "speed": speed,
        "course": course,
        "sat": sat,
    }

    if lat is not None and lon is not None and fix == "1":
        out["google_maps"] = f"https://www.google.com/maps?q={lat},{lon}"
        out["location"] = f"{lat},{lon}"
    return out


def gps_poll_once() -> Dict:
    ensure_gps_on()
    resp = send_at("AT+CGNSINF", max_wait=1.5)
    parsed = parse_cgnsinf(resp)

    if parsed.get("run") != "1":
        gps_power_cycle_sequence(log_it=True)
        resp = send_at("AT+CGNSINF", max_wait=1.5)
        parsed = parse_cgnsinf(resp)

    if parsed.get("fix") == "1" and parsed.get("lat") is not None and parsed.get("lon") is not None:
        db_exec(
            """
            INSERT INTO gps_points
            (id, ts, utc, lat, lon, alt, speed, course, sat, google_maps)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                make_id(),
                now_iso_ph(),
                parsed.get("utc"),
                parsed.get("lat"),
                parsed.get("lon"),
                parsed.get("alt"),
                parsed.get("speed"),
                parsed.get("course"),
                parsed.get("sat"),
                parsed.get("google_maps"),
            )
        )
        db_prune_keep_latest("gps_points", KEEP_GPS_POINTS)

    return parsed


# ===================== Runtime cache =====================
status_cache = {
    "AT": "",
    "CSQ": "",
    "CREG": "",
    "COPS": "",
    "CGPADDR": "",
    "APN": "",
    "AT_MEANING": "",
    "CSQ_MEANING": "",
    "CREG_MEANING": "",
    "POWER": {},
}
gps_latest: Dict = {}
stop_threads = False
last_status_written = {}


# ===================== Threads =====================
def status_thread():
    global status_cache, last_status_written

    while not stop_threads:
        try:
            raw_at = send_at("AT", max_wait=0.8)
            raw_csq = send_at("AT+CSQ", max_wait=1.0)
            raw_creg = send_at("AT+CREG?", max_wait=1.0)
            raw_cops = send_at("AT+COPS?", max_wait=1.2)
            raw_cgpaddr = send_at("AT+CGPADDR=1", max_wait=1.2)
            raw_cgdcont = send_at("AT+CGDCONT?", max_wait=1.2)
            apn = parse_apn_from_cgdcont(raw_cgdcont) or ""
            power = get_throttled_status()

            at_mean = interpret_at(raw_at)
            csq_mean = interpret_csq(raw_csq)
            creg_mean = interpret_creg(raw_creg)

            status_cache = {
                "AT": raw_at,
                "CSQ": raw_csq,
                "CREG": raw_creg,
                "COPS": raw_cops,
                "CGPADDR": raw_cgpaddr,
                "APN": apn,
                "AT_MEANING": at_mean,
                "CSQ_MEANING": csq_mean,
                "CREG_MEANING": creg_mean,
                "POWER": power,
            }

            operator = ""
            m = re.search(r'\+COPS:\s*\d+,\d+,"([^"]+)"', raw_cops)
            if m:
                operator = m.group(1)

            ip = ""
            m = re.search(r"\+CGPADDR:\s*\d+,([0-9.]+)", raw_cgpaddr)
            if m:
                ip = m.group(1)

            current_summary = {
                "at": at_mean,
                "csq": csq_mean,
                "creg": creg_mean,
                "operator": operator,
                "ip": ip,
                "apn": apn,
                "undervoltage_now": power.get("undervoltage_now"),
                "undervoltage_boot": power.get("undervoltage_boot"),
                "throttling_now": power.get("throttling_now"),
                "throttling_boot": power.get("throttling_boot"),
                "throttled_raw": power.get("raw"),
            }

            if current_summary != last_status_written:
                db_exec(
                    """
                    INSERT INTO network_status
                    (id, ts, at, csq, creg, operator, ip, apn,
                     undervoltage_now, undervoltage_boot, throttling_now, throttling_boot, throttled_raw)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        make_id(),
                        now_iso_ph(),
                        at_mean,
                        csq_mean,
                        creg_mean,
                        operator,
                        ip,
                        apn,
                        power.get("undervoltage_now"),
                        power.get("undervoltage_boot"),
                        power.get("throttling_now"),
                        power.get("throttling_boot"),
                        power.get("raw"),
                    )
                )
                db_prune_keep_latest("network_status", KEEP_NETWORK_STATUS)
                last_status_written = current_summary.copy()

        except Exception as e:
            log_activity("ERROR", "STATUS", f"status_thread: {e}")

        time.sleep(STATUS_POLL_SEC)


def gps_thread():
    global gps_latest
    while not stop_threads:
        try:
            gps_latest = gps_poll_once()
        except Exception as e:
            log_activity("ERROR", "GPS", f"gps_thread: {e}")
        time.sleep(GPS_POLL_SEC)


def heartbeat_thread():
    while not stop_threads:
        try:
            hb_touch_online()
        except Exception as e:
            log_activity("WARN", "SYSTEM", f"heartbeat_thread: {e}")
        time.sleep(HEARTBEAT_SEC)


def sms_inbox_thread():
    while not stop_threads:
        try:
            sms_inbox_poll_once()
        except Exception as e:
            log_activity("ERROR", "SMS", f"sms_inbox_thread: {e}")
        time.sleep(SMS_POLL_SEC)


# ===================== Startup/Shutdown =====================
@app.on_event("startup")
def startup():
    global stop_threads
    stop_threads = False

    init_db()
    log_activity("INFO", "SYSTEM", "boot: backend startup")
    hb_touch_online()

    try:
        gps_power_cycle_sequence(log_it=True)
    except Exception as e:
        log_activity("ERROR", "GPS", f"startup_gps_init: {e}")

    try:
        configure_sms_mode()
        log_activity("INFO", "SMS", "startup_sms_init: SMS mode configured")
    except Exception as e:
        log_activity("ERROR", "SMS", f"startup_sms_init: {e}")

    threading.Thread(target=status_thread, daemon=True).start()
    threading.Thread(target=gps_thread, daemon=True).start()
    threading.Thread(target=heartbeat_thread, daemon=True).start()
    threading.Thread(target=sms_inbox_thread, daemon=True).start()


@app.on_event("shutdown")
def shutdown():
    global stop_threads
    stop_threads = True
    log_activity("INFO", "SYSTEM", "power_off: shutdown event triggered")
    try:
        if ser and ser.is_open:
            ser.close()
    except Exception:
        pass


# ===================== Models =====================
class ContactCreate(BaseModel):
    name: str = Field(min_length=1)
    number: str = Field(min_length=10)


class ContactUpdate(BaseModel):
    name: str = Field(min_length=1)
    number: str = Field(min_length=10)


class GroupCreate(BaseModel):
    name: str = Field(min_length=1)


class GroupUpdate(BaseModel):
    members: List[str] = Field(default_factory=list)


class SendSMSRequest(BaseModel):
    to: List[str]
    message: str = Field(min_length=1)


# ===================== API =====================


def mask_to_prefix(mask: str) -> Optional[int]:
    mapping = {
        "255.255.255.252": 30,
        "255.255.255.248": 29,
        "255.255.255.240": 28,
        "255.255.255.0": 24,
        "255.255.0.0": 16,
        "255.0.0.0": 8,
    }
    return mapping.get((mask or "").strip())


def reconnect_mobile_data() -> dict:
    notes = []
    if not Path(QMI_MODEM).exists():
        raise HTTPException(status_code=500, detail=f"QMI modem not found: {QMI_MODEM}")
    if not Path(f"/sys/class/net/{WWAN_IFACE}").exists():
        raise HTTPException(status_code=500, detail=f"WWAN interface not found: {WWAN_IFACE}")

    run_cmd(["ip", "link", "set", WWAN_IFACE, "down"], timeout=5)
    try:
        Path(f"/sys/class/net/{WWAN_IFACE}/qmi/raw_ip").write_text("Y")
    except Exception as e:
        notes.append(f"raw_ip write failed: {e}")
    run_cmd(["ip", "link", "set", WWAN_IFACE, "up"], timeout=5)

    status_before = run_cmd(["qmicli", "-d", QMI_MODEM, "--wds-get-packet-service-status"], timeout=8)
    status_text = status_before["stdout"] or status_before["stderr"]
    if "connected" not in status_text.lower():
        start = run_cmd([
            "qmicli", "-p", "-d", QMI_MODEM,
            '--device-open-net=net-raw-ip|net-no-qos-header',
            "--wds-start-network=apn='internet',ip-type=4",
            '--client-no-release-cid'
        ], timeout=25)
        if not start["ok"]:
            detail = start["stderr"] or start["stdout"] or "Failed to start mobile-data session"
            raise HTTPException(status_code=500, detail=detail)
        notes.append("Started fresh QMI data session")
    else:
        notes.append("Packet service already connected")

    current = run_cmd(["qmicli", "-p", "-d", QMI_MODEM, "--wds-get-current-settings"], timeout=10)
    if not current["ok"]:
        detail = current["stderr"] or current["stdout"] or "Failed to get current QMI settings"
        raise HTTPException(status_code=500, detail=detail)

    parsed = parse_qmi_settings(current["stdout"])
    ip = parsed.get("ipv4")
    mask = parsed.get("subnet_mask")
    gw = parsed.get("gateway")
    dns1 = parsed.get("dns1")
    dns2 = parsed.get("dns2")
    prefix = mask_to_prefix(mask)

    if not ip or not mask or not gw or prefix is None:
        raise HTTPException(status_code=500, detail="Failed to parse IPv4/gateway settings from current QMI session")

    run_cmd(["ip", "addr", "flush", "dev", WWAN_IFACE], timeout=5)
    add_ip = run_cmd(["ip", "addr", "add", f"{ip}/{prefix}", "dev", WWAN_IFACE], timeout=5)
    if not add_ip["ok"]:
        raise HTTPException(status_code=500, detail=add_ip["stderr"] or add_ip["stdout"] or "Failed to assign IPv4 to wwan0")

    route = run_cmd(["ip", "route", "replace", "default", "via", gw, "dev", WWAN_IFACE, "metric", "100"], timeout=5)
    if not route["ok"]:
        raise HTTPException(status_code=500, detail=route["stderr"] or route["stdout"] or "Failed to apply default route")

    try:
        lines = []
        if dns1:
            lines.append(f"nameserver {dns1}")
        if dns2:
            lines.append(f"nameserver {dns2}")
        if lines:
            Path("/etc/resolv.conf").write_text("\n".join(lines) + "\n")
    except Exception as e:
        notes.append(f"Failed to write DNS: {e}")

    after = get_mobile_data_status()
    after["notes"] = (after.get("notes") or []) + notes
    return {
        "ok": True,
        "message": "Mobile data reconnect finished",
        "mobile_data": after,
    }




def tail_text_file(path_str: str, max_lines: int = 50) -> List[str]:
    try:
        p = Path(path_str)
        if not p.exists():
            return []
        with p.open("r", errors="ignore") as f:
            lines = f.readlines()[-max_lines:]
        return [ln.rstrip("\n") for ln in lines]
    except Exception as e:
        return [f"ERROR reading {path_str}: {e}"]



def filtered_kernel_module_logs(max_lines: int = 20) -> List[str]:
    patterns = re.compile(r"sim|ttyUSB|cdc|wwan|qmi|gps|camera|uvc|video|usb", re.I)

    lines = tail_text_file("/var/log/kern.log", max_lines=max_lines * 5)
    if lines:
        filtered = [ln for ln in lines if patterns.search(ln)]
        if filtered:
            return filtered[-max_lines:]

    cmd = run_cmd(["journalctl", "-k", "-n", str(max_lines * 5), "--no-pager"], timeout=8)
    raw = cmd.get("stdout", "") if cmd.get("ok") else ""
    if raw:
        filtered = [ln for ln in raw.splitlines() if patterns.search(ln)]
        if filtered:
            return filtered[-max_lines:]

    cmd = run_cmd(["dmesg", "-T"], timeout=8)
    raw = cmd.get("stdout", "") if cmd.get("ok") else ""
    if raw:
        filtered = [ln for ln in raw.splitlines() if patterns.search(ln)]
        if filtered:
            return filtered[-max_lines:]

    return []


def filtered_ssh_logs(max_lines: int = 50) -> List[str]:
    patterns = re.compile(r"sshd|ssh", re.I)
    lines = tail_text_file("/var/log/auth.log", max_lines=max_lines)
    if lines:
        filtered = [ln for ln in lines if patterns.search(ln)]
        if filtered:
            return filtered[-max_lines:]
    cmd = run_cmd(["journalctl", "-u", "ssh", "-n", str(max_lines), "--no-pager"], timeout=8)
    raw = cmd.get("stdout", "") if cmd.get("ok") else ""
    if not raw:
        cmd = run_cmd(["journalctl", "-u", "sshd", "-n", str(max_lines), "--no-pager"], timeout=8)
        raw = cmd.get("stdout", "") if cmd.get("ok") else cmd.get("stderr", "")
    if not raw:
        return []
    filtered = [ln for ln in raw.splitlines() if patterns.search(ln)]
    return filtered[-max_lines:]


def current_ssh_sessions() -> List[dict]:
    out = []
    cmd = run_cmd(["who"], timeout=4)
    raw = cmd.get("stdout", "") if cmd.get("ok") else ""
    for line in raw.splitlines():
        parts = line.split()
        if len(parts) >= 5:
            out.append({
                "user": parts[0],
                "tty": parts[1],
                "date": f"{parts[2]} {parts[3]}",
                "from": parts[4].strip("()"),
                "raw": line,
            })
        elif len(parts) >= 4:
            out.append({
                "user": parts[0],
                "tty": parts[1],
                "date": f"{parts[2]} {parts[3]}",
                "from": "",
                "raw": line,
            })
    return out


def simple_ssh_access_summary(max_items: int = 20) -> List[str]:
    sessions = current_ssh_sessions()
    lines: List[str] = []
    seen = set()

    # Active sessions first
    for s in sessions:
        user = s.get("user") or "?"
        addr = s.get("from") or "local"
        tty = s.get("tty") or "?"
        item = f"ACTIVE | {user} | {addr} | {tty}"
        if item not in seen:
            lines.append(item)
            seen.add(item)

    # Then recent auth/journal lines, simplified to accepted/failed + address
    raw_lines = filtered_ssh_logs(100)
    accepted_pat = re.compile(r"Accepted\s+\S+\s+for\s+(?P<user>\S+)\s+from\s+(?P<ip>[0-9a-fA-F:.]+)", re.I)
    failed_pat = re.compile(r"Failed\s+\S+\s+for(?:\s+invalid user)?\s+(?P<user>\S+)\s+from\s+(?P<ip>[0-9a-fA-F:.]+)", re.I)
    invalid_pat = re.compile(r"Invalid user\s+(?P<user>\S+)\s+from\s+(?P<ip>[0-9a-fA-F:.]+)", re.I)
    closed_pat = re.compile(r"Disconnected from(?: invalid user)?\s+(?P<user>\S+)?\s*(?P<ip>[0-9a-fA-F:.]+)", re.I)

    for ln in reversed(raw_lines):
        item = None
        m = accepted_pat.search(ln)
        if m:
            item = f"LOGIN OK | {m.group('user')} | {m.group('ip')}"
        else:
            m = failed_pat.search(ln)
            if m:
                item = f"LOGIN FAIL | {m.group('user')} | {m.group('ip')}"
            else:
                m = invalid_pat.search(ln)
                if m:
                    item = f"INVALID USER | {m.group('user')} | {m.group('ip')}"
                else:
                    m = closed_pat.search(ln)
                    if m and m.group('ip'):
                        item = f"DISCONNECTED | {m.group('user') or '?'} | {m.group('ip')}"
        if item and item not in seen:
            lines.append(item)
            seen.add(item)
        if len(lines) >= max_items:
            break

    return lines[:max_items]


def recent_gps_logs(limit: int = 20) -> List[dict]:
    return db_query(
        """
        SELECT ts, utc, lat, lon, speed, course, sat, google_maps
        FROM gps_points
        ORDER BY ts DESC
        LIMIT ?
        """,
        (limit,)
    )


def recent_sms_sent_logs(limit: int = 20) -> List[dict]:
    return db_query(
        """
        SELECT ts, to_number AS number, message, ok, resp
        FROM sms_sent
        ORDER BY ts DESC
        LIMIT ?
        """,
        (limit,)
    )


def recent_sms_inbox_logs(limit: int = 20) -> List[dict]:
    return db_query(
        """
        SELECT ts, from_number, message, raw
        FROM sms_inbox
        ORDER BY ts DESC
        LIMIT ?
        """,
        (limit,)
    )


def recent_activity_logs(limit: int = 50) -> List[dict]:
    return db_query(
        "SELECT ts, level, category, detail FROM activity ORDER BY ts DESC LIMIT ?",
        (limit,)
    )




def present_connected_devices() -> List[str]:
    items: List[str] = []

    if Path(QMI_MODEM).exists():
        items.append(f"SIM7000 QMI modem present: {QMI_MODEM}")
    else:
        items.append(f"SIM7000 QMI modem missing: {QMI_MODEM}")

    if Path(f"/sys/class/net/{WWAN_IFACE}").exists():
        items.append(f"WWAN interface present: {WWAN_IFACE}")
    else:
        items.append(f"WWAN interface missing: {WWAN_IFACE}")

    tty_usbs = sorted(glob.glob("/dev/ttyUSB*"))
    if tty_usbs:
        items.append("USB modem serial ports: " + ", ".join(tty_usbs))
    else:
        items.append("USB modem serial ports: none")

    videos = sorted(glob.glob("/dev/video*"))
    if videos:
        items.append("Video devices: " + ", ".join(videos))
    else:
        items.append("Video devices: none")

    lsusb = run_cmd(["lsusb"], timeout=5)
    if lsusb["ok"] and lsusb["stdout"]:
        lines = [ln.strip() for ln in lsusb["stdout"].splitlines() if ln.strip()]
        lines = lines[-12:]
        items.append("USB devices now connected:")
        items.extend([f"  {ln}" for ln in lines])

    return items

def ongoing_logs_payload() -> dict:
    return {
        "activity": recent_activity_logs(50),
        "gps": recent_gps_logs(20),
        "sms_sent": recent_sms_sent_logs(20),
        "sms_inbox": recent_sms_inbox_logs(20),
        "cameras": list_cameras(),
        "mobile_data": get_mobile_data_status(),
        "present_devices": present_connected_devices(),
        "ssh_sessions": current_ssh_sessions(),
        "ssh_access_summary": simple_ssh_access_summary(20),
        "ssh_log_tail": filtered_ssh_logs(50),
        "kernel_module_log_tail": filtered_kernel_module_logs(20),
    }

@app.get("/mobile_data/status")
def mobile_data_status():
    return get_mobile_data_status()


@app.post("/mobile_data/reconnect")
def mobile_data_reconnect():
    return reconnect_mobile_data()


@app.get("/status")
def get_status():
    groups = db_query("SELECT name FROM groups ORDER BY name ASC")
    hb = hb_get_status()
    mobile = get_mobile_data_status()

    latest_net_rows = db_query(
        """
        SELECT ts, at, csq, creg, operator, ip, apn,
               undervoltage_now, undervoltage_boot, throttling_now, throttling_boot, throttled_raw
        FROM network_status
        ORDER BY ts DESC
        LIMIT 1
        """
    )
    latest_net = latest_net_rows[0] if latest_net_rows else {}

    return {
        "AT": status_cache.get("AT", ""),
        "CSQ": status_cache.get("CSQ", ""),
        "CREG": status_cache.get("CREG", ""),
        "COPS": status_cache.get("COPS", ""),
        "CGPADDR": status_cache.get("CGPADDR", ""),
        "APN": status_cache.get("APN", ""),
        "AT_MEANING": status_cache.get("AT_MEANING", ""),
        "CSQ_MEANING": status_cache.get("CSQ_MEANING", ""),
        "CREG_MEANING": status_cache.get("CREG_MEANING", ""),
        "power": status_cache.get("POWER", {}),
        "contacts_count": db_query("SELECT COUNT(*) AS c FROM contacts")[0]["c"],
        "groups_count": len(groups),
        "groups": [g["name"] for g in groups],
        "server_time": now_dt_ph().strftime("%m/%d/%Y %I:%M %p"),
        "db_backend": _resolve_db_backend(),
        "heartbeat": hb,
        "latest_network_log": latest_net,
        "mobile_data": mobile,
    }


@app.post("/system/poweroff")
def system_poweroff():
    log_activity("WARN", "SYSTEM", "remote_poweroff: requested from GUI")
    threading.Thread(target=delayed_poweroff, daemon=True).start()
    return {
        "ok": True,
        "message": "Poweroff initiated. Wait until the Pi fully stops before unplugging power."
    }


@app.post("/gps/restart")
def gps_restart():
    result = gps_power_cycle_sequence(log_it=True)
    if not result.get("ok"):
        raise HTTPException(500, detail=result.get("error", "Failed to restart GPS"))
    return {
        "ok": True,
        "message": "GPS power-cycled successfully. Wait a bit for satellite fix.",
        "power_state": result.get("power_state"),
    }


@app.get("/contacts")
def contacts_list():
    return {
        "contacts": db_query(
            "SELECT id, name, number, created_at FROM contacts ORDER BY name ASC"
        )
    }


@app.post("/contacts")
def contacts_add(c: ContactCreate):
    number = normalize_number(c.number)
    if not is_valid_ph(number):
        raise HTTPException(400, detail="Number must be 09xxxxxxxxx")

    cid = make_id()
    db_exec(
        "INSERT INTO contacts (id, name, number, created_at) VALUES (?, ?, ?, ?)",
        (cid, c.name.strip(), number, now_iso_ph())
    )
    log_activity("INFO", "SYSTEM", f"contact_add: {c.name.strip()} {number}")
    return {"ok": True, "contact": {"id": cid, "name": c.name.strip(), "number": number}}


@app.put("/contacts/{contact_id}")
def contacts_update(contact_id: str, c: ContactUpdate):
    number = normalize_number(c.number)
    if not is_valid_ph(number):
        raise HTTPException(400, detail="Number must be 09xxxxxxxxx")

    found = db_query("SELECT id FROM contacts WHERE id=?", (contact_id,))
    if not found:
        raise HTTPException(404, detail="Contact not found")

    db_exec(
        "UPDATE contacts SET name=?, number=? WHERE id=?",
        (c.name.strip(), number, contact_id)
    )
    log_activity("INFO", "SYSTEM", f"contact_update: {contact_id} -> {c.name.strip()} {number}")
    return {"ok": True}


@app.delete("/contacts/{contact_id}")
def contacts_delete(contact_id: str):
    found = db_query("SELECT id, name, number FROM contacts WHERE id=?", (contact_id,))
    if not found:
        raise HTTPException(404, detail="Contact not found")

    db_exec("DELETE FROM group_members WHERE contact_id=?", (contact_id,))
    db_exec("DELETE FROM contacts WHERE id=?", (contact_id,))
    log_activity("INFO", "SYSTEM", f"contact_delete: {found[0]['name']} {found[0]['number']}")
    return {"ok": True}


@app.get("/groups")
def groups_list():
    return {"groups": [g["name"] for g in db_query("SELECT name FROM groups ORDER BY name ASC")]}


@app.post("/groups")
def groups_create(g: GroupCreate):
    name = g.name.strip()
    if not name:
        raise HTTPException(400, detail="Group name required")
    exists = db_query("SELECT name FROM groups WHERE name=?", (name,))
    if exists:
        raise HTTPException(400, detail="Group already exists")

    db_exec("INSERT INTO groups (name) VALUES (?)", (name,))
    log_activity("INFO", "SYSTEM", f"group_create: {name}")
    return {"ok": True, "name": name}


@app.get("/groups/{group_name}")
def groups_get(group_name: str):
    exists = db_query("SELECT name FROM groups WHERE name=?", (group_name,))
    if not exists:
        raise HTTPException(404, detail="Group not found")

    members = db_query("SELECT contact_id FROM group_members WHERE group_name=?", (group_name,))
    member_ids = [m["contact_id"] for m in members]

    contacts = []
    if member_ids:
        q = ",".join(["?"] * len(member_ids))
        contacts = db_query(
            f"SELECT id, name, number, created_at FROM contacts WHERE id IN ({q})",
            tuple(member_ids)
        )

    return {"name": group_name, "members": member_ids, "contacts": contacts}


@app.put("/groups/{group_name}")
def groups_update(group_name: str, g: GroupUpdate):
    exists = db_query("SELECT name FROM groups WHERE name=?", (group_name,))
    if not exists:
        raise HTTPException(404, detail="Group not found")

    known_ids = {c["id"] for c in db_query("SELECT id FROM contacts")}
    members = [m for m in g.members if m in known_ids]

    db_exec("DELETE FROM group_members WHERE group_name=?", (group_name,))
    for mid in members:
        db_exec(
            "INSERT INTO group_members (group_name, contact_id) VALUES (?, ?)",
            (group_name, mid)
        )

    log_activity("INFO", "SYSTEM", f"group_update: {group_name} members={len(members)}")
    return {"ok": True, "name": group_name, "members": members}


@app.delete("/groups/{group_name}")
def groups_delete(group_name: str):
    exists = db_query("SELECT name FROM groups WHERE name=?", (group_name,))
    if not exists:
        raise HTTPException(404, detail="Group not found")

    db_exec("DELETE FROM group_members WHERE group_name=?", (group_name,))
    db_exec("DELETE FROM groups WHERE name=?", (group_name,))
    log_activity("INFO", "SYSTEM", f"group_delete: {group_name}")
    return {"ok": True}


@app.post("/send_sms")
def send_sms(req: SendSMSRequest):
    to_list = [normalize_number(n) for n in req.to]
    bad = [n for n in to_list if not is_valid_ph(n)]
    if bad:
        raise HTTPException(400, detail=f"Invalid numbers: {bad}")

    details = []
    for n in to_list:
        sid = make_id()
        try:
            resp = send_sms_single(n, req.message)
            ok = 1 if ("OK" in resp or "+CMGS" in resp) else 0
            db_exec(
                "INSERT INTO sms_sent (id, ts, to_number, message, resp, ok) VALUES (?, ?, ?, ?, ?, ?)",
                (sid, now_iso_ph(), n, req.message, resp, ok)
            )
            details.append({"number": n, "resp": resp})
        except Exception as e:
            db_exec(
                "INSERT INTO sms_sent (id, ts, to_number, message, resp, ok) VALUES (?, ?, ?, ?, ?, ?)",
                (sid, now_iso_ph(), n, req.message, f"ERROR: {e}", 0)
            )
            log_activity("ERROR", "SMS", f"send_sms: {e}")
            details.append({"number": n, "resp": f"ERROR: {e}"})

    db_prune_keep_latest("sms_sent", KEEP_SMS_SENT)
    log_activity("INFO", "SYSTEM", f"sms_send: count={len(to_list)}")
    return {"result": "ok", "count": len(to_list), "details": details}


@app.get("/messages")
def messages_get():
    payload = ongoing_logs_payload()
    return {
        "activity": payload["activity"][:20],
        "sms": payload["sms_sent"][:20],
        "sms_inbox": payload["sms_inbox"][:20],
        "gps": payload["gps"],
        "cameras": payload["cameras"],
        "mobile_data": payload["mobile_data"],
        "ssh_sessions": payload["ssh_sessions"],
        "ssh_access_summary": payload["ssh_access_summary"],
        "ssh_log_tail": payload["ssh_log_tail"],
        "kernel_module_log_tail": payload["kernel_module_log_tail"],
    }


@app.get("/logs/ongoing")
def logs_ongoing():
    return ongoing_logs_payload()


@app.get("/sms_inbox")
def sms_inbox_get():
    rows = db_query(
        """
        SELECT ts, from_number, message, raw
        FROM sms_inbox
        ORDER BY ts DESC
        LIMIT 100
        """
    )
    return {"messages": rows}


@app.get("/gps_latest")
def gps_latest_api():
    if not gps_latest:
        return gps_poll_once()
    return gps_latest


@app.get("/gps_track")
def gps_track_api():
    points = db_query(
        """
        SELECT ts, utc, lat, lon, alt, speed, course, sat
        FROM gps_points
        ORDER BY ts ASC
        LIMIT 2000
        """
    )
    cleaned = []
    for p in points:
        if p.get("lat") is None or p.get("lon") is None:
            continue
        cleaned.append({
            "ts": p["ts"],
            "utc": p["utc"],
            "lat": float(p["lat"]),
            "lon": float(p["lon"]),
            "alt": p.get("alt"),
            "speed": p.get("speed"),
            "course": p.get("course"),
            "sat": p.get("sat"),
        })
    return {"count": len(cleaned), "poll_sec": GPS_POLL_SEC, "max_points": 2000, "points": cleaned}


@app.post("/gps_track/clear")
def gps_track_clear():
    db_exec("DELETE FROM gps_points")
    log_activity("INFO", "SYSTEM", "gps_clear: gps_points cleared")
    return {"ok": True}


@app.get("/cameras")
def cameras_get():
    cams = list_cameras()
    return {
        "count": len(cams),
        "active_streams": sum(1 for c in cams if c.get("streaming")),
        "cameras": cams,
    }


@app.get("/camera/{camera_index}/snapshot.jpg")
def camera_snapshot(camera_index: int):
    try:
        jpg = capture_camera_jpeg(camera_index)
        return Response(content=jpg, media_type="image/jpeg")
    except Exception as e:
        raise HTTPException(500, detail=str(e))


@app.get("/camera/{camera_index}/stream.mjpg")
def camera_stream(camera_index: int):
    try:
        return StreamingResponse(
            mjpeg_stream_generator(camera_index),
            media_type="multipart/x-mixed-replace; boundary=frame"
        )
    except Exception as e:
        raise HTTPException(500, detail=str(e))


def write_multisheet_xlsx(path: Path, sheets: List[tuple[str, List[str], List[dict]]]):
    wb = Workbook()
    first = True
    for title, headers, rows in sheets:
        ws = wb.active if first else wb.create_sheet()
        ws.title = (title or "Sheet")[:31]
        first = False
        ws.append(headers)
        for r in rows:
            out_row = []
            for h in headers:
                v = r.get(h, "")
                if h in ("ts", "created_at", "last_seen_ts", "last_online_ts", "last_offline_ts") and v:
                    out_row.append(excel_ts(v))
                else:
                    if isinstance(v, (dict, list)):
                        v = json.dumps(v, ensure_ascii=False)
                    out_row.append(v)
            ws.append(out_row)
        for col_name in ("ts", "created_at", "last_seen_ts", "last_online_ts", "last_offline_ts"):
            if col_name in headers:
                col = headers.index(col_name) + 1
                for cellcol in ws.iter_cols(min_col=col, max_col=col, min_row=2):
                    for c in cellcol:
                        if isinstance(c.value, datetime):
                            c.number_format = "m/d/yyyy h:mm AM/PM"
    wb.save(path)


@app.get("/export/ongoing_logs.xlsx")
def export_ongoing_logs():
    payload = ongoing_logs_payload()
    out = EXPORT_DIR / f"ongoing_logs_{now_dt_ph().strftime('%Y%m%d_%H%M%S')}.xlsx"

    mobile = payload.get("mobile_data", {}) or {}
    mobile_rows = [{k: (json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v) for k, v in mobile.items() if k != "raw"}]
    if mobile.get("raw"):
        mobile_rows[0]["raw"] = json.dumps(mobile.get("raw"), ensure_ascii=False)
    camera_rows = payload.get("cameras", []) or []
    ssh_rows = payload.get("ssh_sessions", []) or []
    ssh_summary_rows = [{"summary": line} for line in (payload.get("ssh_access_summary", []) or [])]
    present_device_rows = [{"summary": line} for line in (payload.get("present_devices", []) or [])]
    ssh_log_rows = [{"line": line} for line in (payload.get("ssh_log_tail", []) or [])]
    kernel_rows = [{"line": line} for line in (payload.get("kernel_module_log_tail", []) or [])]

    sheets = [
        ("Activity", ["ts", "level", "category", "detail"], payload.get("activity", []) or []),
        ("GPS", ["ts", "utc", "lat", "lon", "speed", "sat", "google_maps"], payload.get("gps", []) or []),
        ("SMS Sent", ["ts", "number", "message", "ok", "resp"], payload.get("sms_sent", []) or []),
        ("SMS Inbox", ["ts", "from_number", "message", "raw"], payload.get("sms_inbox", []) or []),
        ("PresentDevices", ["summary"], present_device_rows),
        ("Cameras", sorted({k for row in camera_rows for k in row.keys()}) or ["device", "exists"], camera_rows),
        ("Mobile Data", sorted(mobile_rows[0].keys()) if mobile_rows else ["connected"], mobile_rows),
        ("SSH Sessions", sorted({k for row in ssh_rows for k in row.keys()}) or ["raw"], ssh_rows),
        ("SSH Log Tail", ["line"], ssh_log_rows),
        ("Kernel Log Tail", ["line"], kernel_rows),
    ]

    write_multisheet_xlsx(out, sheets)
    return FileResponse(out, filename=out.name)


# ===================== Excel Export =====================
def write_xlsx(path: Path, headers: List[str], rows: List[dict]):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)

    for r in rows:
        out_row = []
        for h in headers:
            v = r.get(h, "")
            if h in ("ts", "created_at", "last_seen_ts", "last_online_ts", "last_offline_ts") and v:
                out_row.append(excel_ts(v))
            else:
                out_row.append(v)
        ws.append(out_row)

    for col_name in ("ts", "created_at", "last_seen_ts", "last_online_ts", "last_offline_ts"):
        if col_name in headers:
            col = headers.index(col_name) + 1
            for cellcol in ws.iter_cols(min_col=col, max_col=col, min_row=2):
                for c in cellcol:
                    if isinstance(c.value, datetime):
                        c.number_format = "m/d/yyyy h:mm AM/PM"

    wb.save(path)


@app.get("/export/network_status.xlsx")
def export_network_status():
    rows = db_query(
        """
        SELECT ts, at, csq, creg, operator, ip, apn,
               undervoltage_now, undervoltage_boot, throttling_now, throttling_boot, throttled_raw
        FROM network_status
        ORDER BY ts ASC
        """
    )
    out = EXPORT_DIR / f"network_status_{now_dt_ph().strftime('%Y%m%d_%H%M%S')}.xlsx"
    headers = ["ts", "at", "csq", "creg", "operator", "ip", "apn",
               "undervoltage_now", "undervoltage_boot", "throttling_now", "throttling_boot", "throttled_raw"]
    write_xlsx(out, headers, rows)
    return FileResponse(out, filename=out.name)


@app.get("/export/gps_points.xlsx")
def export_gps_points():
    rows = db_query(
        "SELECT ts, utc, lat, lon, google_maps, sat, speed, course, alt FROM gps_points ORDER BY ts ASC"
    )
    out = EXPORT_DIR / f"gps_points_{now_dt_ph().strftime('%Y%m%d_%H%M%S')}.xlsx"
    headers = ["ts", "utc", "lat", "lon", "google_maps", "sat", "speed", "course", "alt"]
    write_xlsx(out, headers, rows)
    return FileResponse(out, filename=out.name)


@app.get("/export/group_members.xlsx")
def export_group_members():
    rows = db_query("SELECT group_name, contact_id FROM group_members ORDER BY group_name ASC")
    out = EXPORT_DIR / f"group_members_{now_dt_ph().strftime('%Y%m%d_%H%M%S')}.xlsx"
    headers = ["group_name", "contact_id"]
    write_xlsx(out, headers, rows)
    return FileResponse(out, filename=out.name)


@app.get("/export/contacts.xlsx")
def export_contacts():
    rows = db_query("SELECT id, name, number, created_at FROM contacts ORDER BY name ASC")
    out = EXPORT_DIR / f"contacts_{now_dt_ph().strftime('%Y%m%d_%H%M%S')}.xlsx"
    headers = ["id", "name", "number", "created_at"]
    write_xlsx(out, headers, rows)
    return FileResponse(out, filename=out.name)


@app.get("/export/groups.xlsx")
def export_groups():
    rows = db_query("SELECT name FROM groups ORDER BY name ASC")
    out = EXPORT_DIR / f"groups_{now_dt_ph().strftime('%Y%m%d_%H%M%S')}.xlsx"
    headers = ["name"]
    write_xlsx(out, headers, rows)
    return FileResponse(out, filename=out.name)


@app.get("/export/sms_sent.xlsx")
def export_sms_sent():
    rows = db_query("SELECT ts, to_number, message, ok, resp FROM sms_sent ORDER BY ts ASC")
    out = EXPORT_DIR / f"sms_sent_{now_dt_ph().strftime('%Y%m%d_%H%M%S')}.xlsx"
    headers = ["ts", "to_number", "message", "ok", "resp"]
    write_xlsx(out, headers, rows)
    return FileResponse(out, filename=out.name)


@app.get("/export/sms_inbox.xlsx")
def export_sms_inbox():
    rows = db_query("SELECT ts, from_number, message, raw FROM sms_inbox ORDER BY ts ASC")
    out = EXPORT_DIR / f"sms_inbox_{now_dt_ph().strftime('%Y%m%d_%H%M%S')}.xlsx"
    headers = ["ts", "from_number", "message", "raw"]
    write_xlsx(out, headers, rows)
    return FileResponse(out, filename=out.name)


@app.get("/export/heartbeat.xlsx")
def export_heartbeat():
    rows = db_query(
        "SELECT device_id, last_seen_ts, last_online_ts, last_offline_ts FROM heartbeat ORDER BY device_id ASC"
    )
    out = EXPORT_DIR / f"heartbeat_{now_dt_ph().strftime('%Y%m%d_%H%M%S')}.xlsx"
    headers = ["device_id", "last_seen_ts", "last_online_ts", "last_offline_ts"]
    write_xlsx(out, headers, rows)
    return FileResponse(out, filename=out.name)


@app.get("/export/activity.xlsx")
def export_activity():
    rows = db_query("SELECT ts, level, category, detail FROM activity ORDER BY ts ASC")
    out = EXPORT_DIR / f"activity_{now_dt_ph().strftime('%Y%m%d_%H%M%S')}.xlsx"
    headers = ["ts", "level", "category", "detail"]
    write_xlsx(out, headers, rows)
    return FileResponse(out, filename=out.name)


@app.post("/admin/reset_db")
def admin_reset_db():
    global stop_threads, last_status_written, gps_latest

    stop_threads = True
    time.sleep(0.5)

    try:
        if ser and ser.is_open:
            ser.close()
    except Exception:
        pass

    try:
        db_reset_storage()
    except Exception as e:
        raise HTTPException(500, detail=f"Failed to remove DB: {e}")

    init_db()
    last_status_written = {}
    gps_latest = {}
    log_activity("INFO", "SYSTEM", "admin_reset_db: database recreated")
    hb_touch_online()

    try:
        gps_power_cycle_sequence(log_it=True)
    except Exception as e:
        log_activity("ERROR", "GPS", f"reset_gps_init: {e}")

    try:
        configure_sms_mode()
    except Exception as e:
        log_activity("ERROR", "SMS", f"reset_sms_init: {e}")

    stop_threads = False
    threading.Thread(target=status_thread, daemon=True).start()
    threading.Thread(target=gps_thread, daemon=True).start()
    threading.Thread(target=heartbeat_thread, daemon=True).start()
    threading.Thread(target=sms_inbox_thread, daemon=True).start()

    return {"ok": True, "message": "Database reset and recreated."}



